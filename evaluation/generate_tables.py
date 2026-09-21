import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


INPUT_FILE = "evaluation_results.json"
OUTPUT_FILE = "evaluation_tables.xlsx"


def format_sheet(ws):
    # Make headers bold
    for cell in ws[1]:
        cell.font = Font(bold=True)

    # Freeze header row
    ws.freeze_panes = "A2"

    # Automatically adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = min(max_length + 3, 40)


def main():

    # Load JSON
    with open(INPUT_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)

    workbook = Workbook()

    # =========================================================
    # TABLE 1 - Detection Performance
    # =========================================================

    ws = workbook.active
    ws.title = "Detection Performance"

    ws.append([
        "Attack Class",
        "TP",
        "FP",
        "FN",
        "TN",
        "Precision",
        "Recall",
        "F1 Score"
    ])

    for row in data["detection_performance"]["classes"]:
        ws.append([
            row["attack_class"],
            row["tp"],
            row["fp"],
            row["fn"],
            row["tn"],
            row["precision"],
            row["recall"],
            row["f1"]
        ])

    dp = data["detection_performance"]

    ws.append([
        "Macro Average",
        "",
        "",
        "",
        "",
        dp["macro_precision"],
        dp["macro_recall"],
        dp["macro_f1"]
    ])

    format_sheet(ws)

    # =========================================================
    # TABLE 2 - False Positive Analysis
    # =========================================================

    ws = workbook.create_sheet("False Positives")

    ws.append([
        "Metric",
        "Value"
    ])

    fp = data["false_positive_analysis"]

    ws.append(["Benign Cases", fp["n_benign"]])
    ws.append(["Benign Cases Blocked", fp["n_benign_blocked"]])
    ws.append([
        "Aggregate False Positive Rate",
        fp["aggregate_false_positive_rate"]
    ])

    ws.append([])

    ws.append([
        "Detector",
        "False Alarm Cases",
        "False Alarm Rate",
        "Case IDs"
    ])

    for row in fp["per_detector"]:
        ws.append([
            row["detector"],
            row["false_alarm_cases"],
            row["false_alarm_rate"],
            ", ".join(row["cases"])
        ])

    format_sheet(ws)

    # =========================================================
    # TABLE 3 - Evasion Resistance
    # =========================================================

    ws = workbook.create_sheet("Evasion Resistance")

    ws.append([
        "Case ID",
        "Family",
        "Variant",
        "Expected",
        "Caught",
        "Action",
        "Risk Score",
        "Rules"
    ])

    for row in data["evasion_resistance"]["cases"]:
        ws.append([
            row["case_id"],
            row["family"],
            row["variant"],
            row["expected"],
            row["caught_as_expected"],
            row["action"],
            row["risk_score"],
            ", ".join(row["rules"])
        ])

    format_sheet(ws)

    # =========================================================
    # TABLE 4 - Cross Framework
    # =========================================================

    ws = workbook.create_sheet("Cross Framework")

    ws.append([
        "Case ID",
        "Expected Label",
        "Django Status",
        "Flask Status",
        "Django Blocked",
        "Flask Blocked",
        "Same Verdict"
    ])

    for row in data["cross_framework"]["cases"]:
        ws.append([
            row["case_id"],
            row["expected_label"],
            row["django_status"],
            row["flask_status"],
            row["django_blocked"],
            row["flask_blocked"],
            row["same_verdict"]
        ])

    format_sheet(ws)

    # Summary
    ws.append([])
    ws.append([
        "Total Cases",
        data["cross_framework"]["n_cases"]
    ])
    ws.append([
        "Matching Verdicts",
        data["cross_framework"]["matching_verdicts"]
    ])
    ws.append([
        "Consistency Rate",
        data["cross_framework"]["consistency_rate"]
    ])

    # =========================================================
    # SAVE
    # =========================================================

    workbook.save(OUTPUT_FILE)

    print(f"Tables generated successfully: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()